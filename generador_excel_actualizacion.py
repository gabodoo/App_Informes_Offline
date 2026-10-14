import re
from pathlib import Path
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side


def get_base_dir() -> Path:
    import sys
    if getattr(sys, "frozen", False):
        exe_dir = Path(sys.executable).resolve().parent
        if exe_dir.name.lower() == "main" and exe_dir.parent.name.lower() == "dist":
            return exe_dir.parent.parent
        if exe_dir.name.lower() == "dist":
            return exe_dir.parent
        return exe_dir
    return Path(__file__).resolve().parent


class GeneradorExcelActualizacionSuspendido:
    """Clase para la generación del Excel de Actualización: Suspendido a Realizando Estudios."""

    def generar(self, contexto: dict, log_callback=None) -> Path:
        def log(msg):
            if log_callback:
                log_callback(msg)

        rutas_candidatas = [
            get_base_dir() / "plantillas" / "Plantilla _Actualizacion_suspendido_realizando_estudios.xlsx",
            get_base_dir() / "plantillas" / "Plantilla_Actualizacion_suspendido_realizando_estudios.xlsx",
            get_base_dir() / "dist" / "main" / "plantillas" / "Plantilla _Actualizacion_suspendido_realizando_estudios.xlsx",
            Path(__file__).resolve().parent / "plantillas" / "Plantilla _Actualizacion_suspendido_realizando_estudios.xlsx",
        ]
        ruta_plantilla = next((r for r in rutas_candidatas if r.is_file()), rutas_candidatas[0])
        if not ruta_plantilla.exists():
            raise FileNotFoundError(f"Falta plantilla Excel: {ruta_plantilla}")

        log("Cargando plantilla Excel de actualización...")
        wb = openpyxl.load_workbook(ruta_plantilla)
        ws = wb.active

        becarios = contexto.get("becarios", [])
        
        # Fila base donde comienzan los datos (Fila 1 = Título, Fila 2 = Vacía, Fila 3 = Cabecera azul)
        fila_base = 4
        
        # Limpiar filas de datos previas a partir de fila 4
        ultima_fila = ws.max_row
        if ultima_fila >= fila_base:
            ws.delete_rows(fila_base, ultima_fila - fila_base + 1)

        # Estilos especificados por el usuario
        # Arial 11, color negro, mayúsculas, centrado horizontal y vertical
        fuente_datos = Font(name="Arial", size=11, color="000000", bold=False)
        alineacion_centrada = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        borde_fino = Border(
            left=Side(style="thin", color="000000"),
            right=Side(style="thin", color="000000"),
            top=Side(style="thin", color="000000"),
            bottom=Side(style="thin", color="000000")
        )

        # Plomo bajito para columnas K a R (columnas 11 a 18)
        plomo_bajito_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")

        for i, b in enumerate(becarios):
            fila_actual = fila_base + i
            ws.row_dimensions[fila_actual].height = 40.0

            # Extracción limpia de datos en mayúsculas
            exp_val = b.get("expediente", "")
            if str(exp_val).isdigit():
                exp_val = int(exp_val)
            else:
                exp_val = str(exp_val).strip().upper()

            dni_val = str(b.get("dni", "")).strip().upper()
            becario_val = str(b.get("becario", "")).strip().upper()
            
            # BECA: programa - año_convocatoria
            beca_val = str(b.get("beca_excel") or b.get("beca") or "").strip().upper()
            
            inst_val = str(b.get("institucion", "")).strip().upper()
            fin_val = str(b.get("fin", "")).strip().upper()

            # Columnas A a J (1 a 10)
            valores_datos = [
                i + 1,                  # A (1): N°
                exp_val,                # B (2): Expediente
                dni_val,                # C (3): DNI
                becario_val,            # D (4): BECARIO
                beca_val,               # E (5): BECA
                inst_val,               # F (6): INSTITUCIÓN
                fin_val,                # G (7): FECHA FIN SEGÚN SIBEC
                "ACTIVO",               # H (8): NUEVO ESTADO
                "REGULAR",              # I (9): NUEVA CONDICIÓN
                "REALIZANDO ESTUDIOS",  # J (10): NUEVA SITUACIÓN
            ]

            for c, val in enumerate(valores_datos, start=1):
                celda = ws.cell(row=fila_actual, column=c)
                celda.value = val
                celda.font = fuente_datos
                celda.alignment = alineacion_centrada
                celda.border = borde_fino

            # Columnas K a R (11 a 18): Vacías sombreadas con plomo bajito
            for c in range(11, 19):
                celda = ws.cell(row=fila_actual, column=c)
                celda.value = None
                celda.font = fuente_datos
                celda.alignment = alineacion_centrada
                celda.border = borde_fino
                celda.fill = plomo_bajito_fill

        nro_informe = str(contexto.get("nro_informe_word", "")).replace("/", "-").replace("\\", "-")
        nombre_salida = f"Actualizacion_Estado_Condicion_{nro_informe}.xlsx"
        ruta_salida = get_base_dir() / "Informes_Generados" / nombre_salida
        ruta_salida.parent.mkdir(parents=True, exist_ok=True)

        if ruta_salida.exists():
            try:
                with open(ruta_salida, "a+b"):
                    pass
            except (PermissionError, IOError):
                ts = datetime.now().strftime("%H%M%S")
                nombre_salida = f"Actualizacion_Estado_Condicion_{nro_informe}_{ts}.xlsx"
                ruta_salida = get_base_dir() / "Informes_Generados" / nombre_salida

        log(f"Guardando Excel Actualización en: {ruta_salida.name}")
        wb.save(ruta_salida)
        return ruta_salida
