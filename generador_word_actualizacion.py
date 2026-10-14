import re
from pathlib import Path
from datetime import datetime
import docx
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docxtpl import DocxTemplate
from copy import copy
from docx.oxml import parse_xml


def get_base_dir() -> Path:
    import sys
    if getattr(sys, "frozen", False):
        exe_dir = Path(sys.executable).resolve().parent
        if exe_dir.name.lower() == "main" and exe_dir.parent.name.lower() == "dist":
            return exe_dir.parent.parent
        if exe_dir.name.lower() == "dist":
            return exe_dir.parent
        return exe_dir
    return Path(__file__).resolve().parent


class GeneradorWordActualizacionSuspendido:
    """Clase robusta para la generación del Informe Word de Actualización."""

    MESES = ["", "enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]

    def generar(self, contexto: dict, log_callback=None) -> Path:
        def log(msg):
            if log_callback:
                log_callback(msg)

        rutas_candidatas = [
            get_base_dir() / "plantillas" / "Plantilla_Informe_suspendido_realizando_estudios.docx",
            get_base_dir() / "dist" / "main" / "plantillas" / "Plantilla_Informe_suspendido_realizando_estudios.docx",
            Path(__file__).resolve().parent / "plantillas" / "Plantilla_Informe_suspendido_realizando_estudios.docx",
        ]
        ruta_plantilla = next((r for r in rutas_candidatas if r.is_file()), rutas_candidatas[0])
        if not ruta_plantilla.exists():
            raise FileNotFoundError(f"Falta plantilla Word: {ruta_plantilla}")

        log(f"Cargando plantilla Word: {ruta_plantilla.name}")

        nro_informe = str(contexto.get("NUMERO_INFORME") or contexto.get("nro_informe_word") or "").strip()
        becarios = contexto.get("becarios", [])
        info_succor = contexto.get("informe_succor", {})
        resoluciones = contexto.get("RESOLUCIONES") or contexto.get("resoluciones", "")

        n_becarios = len(becarios)
        str_becarios = contexto.get("CANTIDAD_BECARIOS") or (f"{n_becarios} becarios" if n_becarios > 1 else "1 becario")

        str_becas = contexto.get("BECA_Y_CONVOCATORIA", "Beca 18 - 2022")
        str_inst = contexto.get("INSTITUCION", "la institución")

        num_succor = contexto.get("REFERENCIA") or info_succor.get("numero_informe", "Informe N° S/N-2026-MINEDU/VMGI-PRONABEC-DICONCI-SUCCOR-LIMA")
        sigedo_succor = contexto.get("NUMERO_SIGEDO") or info_succor.get("numero_sigedo", "")
        texto_21_succor = contexto.get("TEXTO_SUCCOR") or info_succor.get("texto_numeral_2_1", "")

        hoy = datetime.now()
        fecha_str = contexto.get("FECHA_ACTUAL_TEXTO") or f"{hoy.day} de {self.MESES[hoy.month]} de {hoy.year}"

        asunto_texto = contexto.get("ASUNTO") or f"Modificación de Estado, Condición y Situación de {str_becarios} de la {str_becas} de la {str_inst}"

        # Preparar contexto limpio y expandido para docxtpl (Paso 1)
        ctx_tpl = {
            "NUMERO_INFORME": nro_informe,
            "NUMERO_INFORME_GENERAR": nro_informe,
            "CANTIDAD_BECARIOS": str_becarios,
            "BECA_Y_CONVOCATORIA": str_becas,
            "INSTITUCION": str_inst,
            "REFERENCIA": num_succor,
            "NOMBRE_INFORME_SUCCOR": num_succor,
            "NUMERO_SIGEDO": sigedo_succor,
            "FECHA_ACTUAL_TEXTO": fecha_str,
            "TEXTO_SUCCOR": texto_21_succor,
            "RESOLUCIONES": resoluciones,
            "ASUNTO": asunto_texto,
        }
        # Añadir variantes mayúsculas/minúsculas
        for k in list(ctx_tpl.keys()):
            val = str(ctx_tpl[k])
            ctx_tpl[k.lower()] = val
            ctx_tpl[k.upper()] = val

        nombre_salida = f"Informe_Actualizacion_{nro_informe}.docx"
        ruta_salida = get_base_dir() / "Informes_Generados" / nombre_salida
        ruta_salida.parent.mkdir(parents=True, exist_ok=True)

        if ruta_salida.exists():
            try:
                with open(ruta_salida, "a+b"):
                    pass
            except (PermissionError, IOError):
                ts = datetime.now().strftime("%H%M%S")
                nombre_salida = f"Informe_Actualizacion_{nro_informe}_{ts}.docx"
                ruta_salida = get_base_dir() / "Informes_Generados" / nombre_salida

        # PASO 1: Renderizado con docxtpl
        try:
            tpl = DocxTemplate(str(ruta_plantilla))
            tpl.render(ctx_tpl)
            tpl.save(str(ruta_salida))
            log("Paso 1 (docxtpl): Renderizado de marcadores completado.")
        except Exception as e:
            log(f"Aviso docxtpl: {e}. Se procederá con reemplazo directo python-docx.")
            import shutil
            shutil.copy2(ruta_plantilla, ruta_salida)

        # PASO 2: Post-procesamiento estricto con python-docx (Formato de títulos, tablas, fuentes)
        doc = Document(ruta_salida)

        # 1. Formato estricto del Título
        # Arial 11 pt, Negrita, Subrayado, Centrado
        for p in doc.paragraphs[:5]:
            if "INFORME N" in p.text.upper() and "MINEDU/VMGI" in p.text.upper():
                p.text = f"INFORME N° {nro_informe}-2026-MINEDU/VMGI-PRONABEC-DIBEC-SUS"
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(11)
                    r.bold = True
                    r.underline = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                break

        # 2. Formato estricto de Cabecera (Tabla 0: Asunto, Referencia, Fecha, SIGEDO)
        if len(doc.tables) > 0:
            t0 = doc.tables[0]
            # Fila 2: ASUNTO
            if len(t0.rows) > 2 and len(t0.rows[2].cells) > 2:
                p_asunto = t0.rows[2].cells[2].paragraphs[0]
                p_asunto.text = asunto_texto
                for r in p_asunto.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(11)
                p_asunto.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            # Fila 3: REFERENCIA
            if len(t0.rows) > 3 and len(t0.rows[3].cells) > 2:
                p_ref = t0.rows[3].cells[2].paragraphs[0]
                p_ref.text = num_succor
                for r in p_ref.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(11)
                p_ref.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            # Fila 4: FECHA y SIGEDO
            if len(t0.rows) > 4 and len(t0.rows[4].cells) > 2:
                p_fecha = t0.rows[4].cells[2].paragraphs[0]
                p_fecha.text = f"San Isidro, {fecha_str}"
                for r in p_fecha.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(11)
                p_fecha.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

                if len(t0.rows[4].cells[2].paragraphs) > 1:
                    p_sig = t0.rows[4].cells[2].paragraphs[1]
                else:
                    p_sig = t0.rows[4].cells[2].add_paragraph()
                p_sig.text = f"SIGEDO: {sigedo_succor}"
                for r in p_sig.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(11)
                p_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        # 3. Reemplazo profundo de marcadores residuales en todo el documento
        reemplazos_texto = [
            ("{{ NUMERO_INFORME }}", nro_informe),
            ("{{NUMERO_INFORME}}", nro_informe),
            ("{{ CANTIDAD_BECARIOS }}", str_becarios),
            ("{{CANTIDAD_BECARIOS}}", str_becarios),
            ("{{ BECA_Y_CONVOCATORIA }}", str_becas),
            ("{{BECA_Y_CONVOCATORIA}}", str_becas),
            ("{{ INSTITUCION }}", str_inst),
            ("{{INSTITUCION}}", str_inst),
            ("{{ REFERENCIA }}", num_succor),
            ("{{REFERENCIA}}", num_succor),
            ("{{ NOMBRE_INFORME_SUCCOR }}", num_succor),
            ("{{NOMBRE_INFORME_SUCCOR}}", num_succor),
            ("{{ NUMERO_SIGEDO }}", sigedo_succor),
            ("{{NUMERO_SIGEDO}}", sigedo_succor),
            ("{{ FECHA_ACTUAL_TEXTO }}", fecha_str),
            ("{{FECHA_ACTUAL_TEXTO}}", fecha_str),
            ("{{ TEXTO_SUCCOR }}", texto_21_succor),
            ("{{TEXTO_SUCCOR}}", texto_21_succor),
            ("{{ RESOLUCIONES }}", resoluciones),
            ("{{RESOLUCIONES}}", resoluciones),
            ("{{ ASUNTO }}", asunto_texto),
            ("{{ASUNTO}}", asunto_texto),
            ("4 becarios de la Beca 18 - Convocatorias 2021 y 2022 de la Universidad Peruana de Ciencias Aplicadas", f"{str_becarios} de la {str_becas} de la {str_inst}"),
            ("4 becarios de la Beca 18 - Convocatorias 2021 y 2022", f"{str_becarios} de la {str_becas}"),
            ("73928-2026", sigedo_succor),
            ("Informe N° 5357-2026-MINEDU/VMGI-PRONABEC-DICONCI-SUCCOR-LIMA", num_succor),
            ("Informe N°  5357-2026-MINEDU/VMGI-PRONABEC-DICONCI-SUCCOR-LIMA", num_succor),
        ]

        def procesar_parrafo(p):
            if not p.text:
                return
            modificado = False
            texto = p.text
            for tag, val in reemplazos_texto:
                if tag in texto:
                    texto = texto.replace(tag, str(val))
                    modificado = True
            if modificado:
                p.text = texto

        for p in doc.paragraphs:
            procesar_parrafo(p)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        procesar_parrafo(p)

        # 4. Formato estricto de párrafos clave (2.1, 2.4, 2.5, 2.7, 3.1)
        def actualizar_parrafo(p, nuevo_texto):
            p.text = nuevo_texto
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for r in p.runs:
                r.font.name = "Arial"
                r.font.size = Pt(11)

        for i, p in enumerate(doc.paragraphs):
            texto = p.text.strip()
            if not texto:
                continue

            # 2.1
            if "solicita la actualización del estado, condición y situación" in texto and "conforme al siguiente detalle" in texto:
                actualizar_parrafo(p, f"Mediante {num_succor}, la Subdirección de Coordinación Nacional y Cooperación Regional Lima solicita la actualización del estado, condición y situación de {str_becarios} de la {str_becas}, conforme al siguiente detalle:")

            # 2.4
            elif "De la revisión de los documentos contenidos en el expediente del sigedo" in texto or "expediente del sigedo" in texto:
                actualizar_parrafo(p, f"De la revisión de los documentos contenidos en el expediente del sigedo N° {sigedo_succor} remitido por la Subdirección de Coordinación Nacional y Cooperación Regional Lima (SUCCOR Lima), se verifica que {str_becarios} de la {str_becas} de la {str_inst} tiene el periodo de suspensión de la beca autorizado mediante {resoluciones}.")

            # 2.5
            elif "han culminado su periodo de suspensión de la beca" in texto or "ha culminado su periodo de suspensión de la beca" in texto:
                verbo = "han culminado" if n_becarios > 1 else "ha culminado"
                pronombre = "les correspondería" if n_becarios > 1 else "le correspondería"
                actualizar_parrafo(p, f"Asimismo, se verifica que {str_becarios} {verbo} su periodo de suspensión de la beca, por lo tanto, {pronombre} reiniciar sus estudios.")

            # 2.7
            elif "se sustenta la necesidad de modificar el estado" in texto:
                actualizar_parrafo(p, f"En consecuencia, se sustenta la necesidad de modificar el estado, condición y situación de {str_becarios}, de acuerdo al siguiente detalle:")

            # 3.1 (Conclusiones)
            elif ("solicita la actualización del estado, condición y situación" in texto and "conforme al siguiente detalle" not in texto) or (i >= 58 and "solicita la actualización del estado" in texto):
                actualizar_parrafo(p, f"La Subdirección de Coordinación Nacional y Cooperación Regional Lima mediante {num_succor} solicita la actualización del estado, condición y situación de {str_becarios} de la {str_becas} de la {str_inst}.")

        # 5. Llenado dinámico de Cuadro N° 01 y Cuadro N° 02
        tabla_01 = None
        tabla_02 = None

        for table in doc.tables[1:]:  # Omitir tabla de cabecera
            if len(table.rows) > 0:
                header_text = " ".join([c.text for c in table.rows[0].cells]).upper()
                if "DNI" in header_text:
                    if "NUEVO ESTADO" in header_text or "NUEVA CONDICION" in header_text or "NUEVA SITUACION" in header_text:
                        tabla_02 = table
                    else:
                        if not tabla_01:
                            tabla_01 = table
                        elif not tabla_02:
                            tabla_02 = table

        if not tabla_01 and len(doc.tables) >= 2:
            tabla_01 = doc.tables[1]
        if not tabla_02 and len(doc.tables) >= 3:
            tabla_02 = doc.tables[2]

        def aplicar_bordes_celda(cell):
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = tcPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcBorders')
            if tcBorders is None:
                tcBorders_xml = parse_xml(
                    r'<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                    r'<w:top w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
                    r'<w:left w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
                    r'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
                    r'<w:right w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
                    r'</w:tcBorders>'
                )
                tcPr.append(tcBorders_xml)

        def set_cell_text(cell, text, align_center=True, bold=False):
            cell.text = str(text).strip()
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            aplicar_bordes_celda(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.0
                if align_center:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(8)
                    r.bold = bold

        def formatear_encabezado_tabla(t):
            if len(t.rows) > 0:
                for cell in t.rows[0].cells:
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    aplicar_bordes_celda(cell)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_before = Pt(3)
                        p.paragraph_format.space_after = Pt(3)
                        p.paragraph_format.line_spacing = 1.0
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        for r in p.runs:
                            r.font.name = "Arial"
                            r.font.size = Pt(8)
                            r.bold = True

        # Rellenar Cuadro N° 01 (Estado actual según padrón)
        if tabla_01:
            formatear_encabezado_tabla(tabla_01)
            if len(becarios) > 0:
                while len(tabla_01.rows) > 1:
                    tabla_01._element.remove(tabla_01.rows[-1]._element)

                for b in becarios:
                    row = tabla_01.add_row()
                    set_cell_text(row.cells[0], b.get("dni", ""), align_center=True, bold=False)
                    set_cell_text(row.cells[1], b.get("becario", ""), align_center=True, bold=False)
                    set_cell_text(row.cells[2], b.get("beca", ""), align_center=True, bold=False)
                    set_cell_text(row.cells[3], b.get("institucion", ""), align_center=True, bold=False)
                    set_cell_text(row.cells[4], b.get("estado", "ACTIVO"), align_center=True, bold=False)
                    set_cell_text(row.cells[5], b.get("condicion", "SUSPENDIDO"), align_center=True, bold=False)
                    set_cell_text(row.cells[6], b.get("situacion", "SUSPENSIÓN"), align_center=True, bold=False)

        # Rellenar Cuadro N° 02 (Nuevo estado regular y realizando estudios)
        if tabla_02:
            formatear_encabezado_tabla(tabla_02)
            if len(becarios) > 0:
                while len(tabla_02.rows) > 1:
                    tabla_02._element.remove(tabla_02.rows[-1]._element)

                for b in becarios:
                    row = tabla_02.add_row()
                    set_cell_text(row.cells[0], b.get("dni", ""), align_center=True, bold=False)
                    set_cell_text(row.cells[1], b.get("becario", ""), align_center=True, bold=False)
                    set_cell_text(row.cells[2], b.get("beca", ""), align_center=True, bold=False)
                    set_cell_text(row.cells[3], b.get("institucion", ""), align_center=True, bold=False)
                    set_cell_text(row.cells[4], "ACTIVO", align_center=True, bold=False)
                    set_cell_text(row.cells[5], "REGULAR", align_center=True, bold=False)
                    set_cell_text(row.cells[6], "REALIZANDO ESTUDIOS", align_center=True, bold=False)

        # 6. Siglas al final
        siglas_texto = "ASV/jdbh/jgds/svrq"
        encontrado_siglas = False
        for p in doc.paragraphs[-5:]:
            if "ASV" in p.text or "jdbh" in p.text:
                p.text = siglas_texto
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(10)
                encontrado_siglas = True
                break

        if not encontrado_siglas:
            p_siglas = None
            for p in reversed(doc.paragraphs[-3:]):
                if not p.text.strip():
                    p_siglas = p
                    break
            if not p_siglas:
                p_siglas = doc.add_paragraph()

            p_siglas.text = siglas_texto
            p_siglas.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_siglas.paragraph_format.space_before = Pt(12)
            p_siglas.paragraph_format.space_after = Pt(0)
            for r in p_siglas.runs:
                r.font.name = "Arial"
                r.font.size = Pt(10)

        log(f"Guardando Informe Word en: {ruta_salida.name}")
        doc.save(str(ruta_salida))
        return ruta_salida
